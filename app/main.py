from __future__ import annotations

import calendar
import hashlib
import json
import mimetypes
import os
import re
import secrets
import shutil
import sqlite3
import urllib.request
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - 运行环境未安装 python-docx 时给出友好提示
    DocxDocument = None

BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / "lan_project_manager.db"
STORAGE_DIR = BASE_DIR / "storage"
TEMPLATE_FILE_DIR = STORAGE_DIR / "templates"
DOCUMENT_FILE_DIR = STORAGE_DIR / "documents"
DOCUMENT_VERSION_DIR = STORAGE_DIR / "document_versions"
TASK_IMAGE_DIR = STORAGE_DIR / "task_images"
PROJECT_FILE_DIR = STORAGE_DIR / "project_files"

for folder in [TEMPLATE_FILE_DIR, DOCUMENT_FILE_DIR, DOCUMENT_VERSION_DIR, TASK_IMAGE_DIR, PROJECT_FILE_DIR]:
    folder.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="局域网小组项目工作进度综合管理系统")
app.add_middleware(SessionMiddleware, secret_key=os.getenv("SESSION_SECRET", secrets.token_hex(32)))
app.mount("/static", StaticFiles(directory=BASE_DIR / "app" / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "app" / "templates")


@contextmanager
def db_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
        conn.commit()
    finally:
        conn.close()


def password_hash(password: str, salt: str | None = None) -> str:
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 120_000)
    return f"{salt}${digest.hex()}"


def verify_password(password: str, stored: str) -> bool:
    try:
        salt, digest = stored.split("$", 1)
    except ValueError:
        return False
    return secrets.compare_digest(password_hash(password, salt).split("$", 1)[1], digest)


def init_db() -> None:
    with db_conn() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                display_name TEXT NOT NULL,
                role TEXT NOT NULL CHECK(role IN ('admin','teacher','student')),
                password_hash TEXT NOT NULL,
                student_id TEXT UNIQUE,
                phone TEXT DEFAULT '',
                email TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT DEFAULT ''
            );

            CREATE TABLE IF NOT EXISTS signins (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                signin_date TEXT NOT NULL,
                created_at TEXT NOT NULL,
                UNIQUE(user_id, signin_date),
                FOREIGN KEY(user_id) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                description TEXT DEFAULT '',
                parent_id INTEGER,
                assigned_to INTEGER,
                target_words INTEGER DEFAULT 0,
                content TEXT DEFAULT '',
                status TEXT DEFAULT '未完成',
                due_date TEXT,
                created_by INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(parent_id) REFERENCES tasks(id),
                FOREIGN KEY(assigned_to) REFERENCES users(id),
                FOREIGN KEY(created_by) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS template_files (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                original_name TEXT NOT NULL,
                stored_name TEXT NOT NULL,
                uploaded_by INTEGER NOT NULL,
                uploaded_at TEXT NOT NULL,
                description TEXT DEFAULT '',
                FOREIGN KEY(uploaded_by) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS personal_tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                owner_id INTEGER NOT NULL,
                title TEXT NOT NULL,
                description TEXT DEFAULT '',
                frequency TEXT NOT NULL DEFAULT 'daily' CHECK(frequency IN ('daily','weekly','once')),
                start_date TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT '未完成' CHECK(status IN ('未完成','已完成')),
                due_date TEXT,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                completed_at TEXT,
                FOREIGN KEY(owner_id) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS document_files (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                original_name TEXT NOT NULL,
                stored_name TEXT NOT NULL,
                file_type TEXT NOT NULL,
                task_id INTEGER,
                extracted_text TEXT DEFAULT '',
                created_by INTEGER NOT NULL,
                updated_by INTEGER,
                uploaded_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                description TEXT DEFAULT '',
                current_version INTEGER DEFAULT 1,
                FOREIGN KEY(task_id) REFERENCES tasks(id),
                FOREIGN KEY(created_by) REFERENCES users(id),
                FOREIGN KEY(updated_by) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS document_versions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                document_id INTEGER NOT NULL,
                version_no INTEGER NOT NULL,
                stored_name TEXT NOT NULL,
                saved_by INTEGER NOT NULL,
                saved_at TEXT NOT NULL,
                word_count INTEGER DEFAULT 0,
                note TEXT DEFAULT '',
                FOREIGN KEY(document_id) REFERENCES document_files(id),
                FOREIGN KEY(saved_by) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS task_images (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                original_name TEXT NOT NULL,
                stored_name TEXT NOT NULL,
                mime_type TEXT DEFAULT '',
                caption TEXT DEFAULT '',
                uploaded_by INTEGER NOT NULL,
                uploaded_at TEXT NOT NULL,
                FOREIGN KEY(task_id) REFERENCES tasks(id),
                FOREIGN KEY(uploaded_by) REFERENCES users(id)
            );
            

            CREATE TABLE IF NOT EXISTS task_progress_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                record_type TEXT NOT NULL CHECK(record_type IN ('assigned','personal')),
                task_id INTEGER NOT NULL,
                owner_id INTEGER NOT NULL,
                block_index INTEGER NOT NULL,
                block_start TEXT NOT NULL,
                block_end TEXT NOT NULL,
                note TEXT DEFAULT '',
                original_name TEXT DEFAULT '',
                stored_name TEXT DEFAULT '',
                mime_type TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                FOREIGN KEY(owner_id) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS system_settings (
                setting_key TEXT PRIMARY KEY,
                setting_value TEXT NOT NULL DEFAULT ''
            );

            CREATE TABLE IF NOT EXISTS admin_creation_requests (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                requested_by INTEGER NOT NULL,
                username TEXT UNIQUE NOT NULL,
                display_name TEXT NOT NULL,
                password_hash TEXT NOT NULL,
                phone TEXT DEFAULT '',
                email TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT '待审批',
                required_approvals INTEGER NOT NULL DEFAULT 2,
                created_at TEXT NOT NULL,
                resolved_at TEXT DEFAULT '',
                created_user_id INTEGER,
                FOREIGN KEY(requested_by) REFERENCES users(id),
                FOREIGN KEY(created_user_id) REFERENCES users(id)
            );

            CREATE TABLE IF NOT EXISTS admin_creation_approvals (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                request_id INTEGER NOT NULL,
                admin_id INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                UNIQUE(request_id, admin_id),
                FOREIGN KEY(request_id) REFERENCES admin_creation_requests(id),
                FOREIGN KEY(admin_id) REFERENCES users(id)
            );
            """
        )

        # 兼容旧版数据库：为正式账号体系补充资料字段。
        existing_columns = {row[1] for row in conn.execute("PRAGMA table_info(users)").fetchall()}
        if "student_id" not in existing_columns:
            conn.execute("ALTER TABLE users ADD COLUMN student_id TEXT")
        if "phone" not in existing_columns:
            conn.execute("ALTER TABLE users ADD COLUMN phone TEXT DEFAULT ''")
        if "email" not in existing_columns:
            conn.execute("ALTER TABLE users ADD COLUMN email TEXT DEFAULT ''")
        if "updated_at" not in existing_columns:
            conn.execute("ALTER TABLE users ADD COLUMN updated_at TEXT DEFAULT ''")
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_users_student_id_unique ON users(student_id) WHERE student_id IS NOT NULL AND student_id <> ''")
        conn.execute("UPDATE users SET updated_at=created_at WHERE updated_at IS NULL OR updated_at=''")

        doc_columns = {row[1] for row in conn.execute("PRAGMA table_info(document_files)").fetchall()}
        if "ai_summary" not in doc_columns:
            conn.execute("ALTER TABLE document_files ADD COLUMN ai_summary TEXT DEFAULT ''")
        if "ai_summary_updated_at" not in doc_columns:
            conn.execute("ALTER TABLE document_files ADD COLUMN ai_summary_updated_at TEXT DEFAULT ''")
        if "ai_summary_provider" not in doc_columns:
            conn.execute("ALTER TABLE document_files ADD COLUMN ai_summary_provider TEXT DEFAULT ''")
        if "deleted_at" not in doc_columns:
            conn.execute("ALTER TABLE document_files ADD COLUMN deleted_at TEXT DEFAULT ''")
        if "deleted_by" not in doc_columns:
            conn.execute("ALTER TABLE document_files ADD COLUMN deleted_by INTEGER")

        personal_columns = {row[1] for row in conn.execute("PRAGMA table_info(personal_tasks)").fetchall()}
        new_recurrence_columns = False
        if "start_date" not in personal_columns:
            conn.execute("ALTER TABLE personal_tasks ADD COLUMN start_date TEXT DEFAULT ''")
            conn.execute("UPDATE personal_tasks SET start_date=substr(created_at, 1, 10) WHERE start_date IS NULL OR start_date=''")
        if "is_template" not in personal_columns:
            conn.execute("ALTER TABLE personal_tasks ADD COLUMN is_template INTEGER DEFAULT 0")
            new_recurrence_columns = True
        if "auto_renew" not in personal_columns:
            conn.execute("ALTER TABLE personal_tasks ADD COLUMN auto_renew INTEGER DEFAULT 0")
            new_recurrence_columns = True
        if "recurrence_parent_id" not in personal_columns:
            conn.execute("ALTER TABLE personal_tasks ADD COLUMN recurrence_parent_id INTEGER")
        if "occurrence_key" not in personal_columns:
            conn.execute("ALTER TABLE personal_tasks ADD COLUMN occurrence_key TEXT DEFAULT ''")
        if new_recurrence_columns:
            conn.execute("""
                UPDATE personal_tasks
                SET is_template=1, auto_renew=1
                WHERE frequency IN ('daily','weekly') AND recurrence_parent_id IS NULL
            """)
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_personal_occurrence_unique ON personal_tasks(recurrence_parent_id, occurrence_key) WHERE recurrence_parent_id IS NOT NULL AND occurrence_key IS NOT NULL AND occurrence_key <> ''")

        user_count = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        if user_count == 0:
            now = datetime.now().isoformat(timespec="seconds")
            users = [
                ("Kuiruru", "Kuiruru", "admin", password_hash("Zxcv？6787"), None, "", "", now, now),
                ("AsterN7q", "AsterN7q", "admin", password_hash("M8!vR2#qL7@x"), None, "", "", now, now),
                ("VantaC4x", "VantaC4x", "admin", password_hash("Q6@pN9!sK3#t"), None, "", "", now, now),
                ("teacher", "指导老师", "teacher", password_hash("teacher123"), None, "", "", now, now),
                ("student1", "学生一", "student", password_hash("student123"), None, "", "", now, now),
                ("student2", "学生二", "student", password_hash("student123"), None, "", "", now, now),
            ]
            conn.executemany(
                """
                INSERT INTO users(username, display_name, role, password_hash, student_id, phone, email, created_at, updated_at)
                VALUES(?,?,?,?,?,?,?,?,?)
                """,
                users,
            )

            teacher_id = conn.execute("SELECT id FROM users WHERE username='teacher'").fetchone()[0]
            s1 = conn.execute("SELECT id FROM users WHERE username='student1'").fetchone()[0]
            s2 = conn.execute("SELECT id FROM users WHERE username='student2'").fetchone()[0]

            conn.execute(
                """
                INSERT INTO tasks(title, description, parent_id, assigned_to, target_words, created_by, created_at, updated_at)
                VALUES(?,?,?,?,?,?,?,?)
                """,
                ("太空能源技术概论（示例书稿）", "来自上传样例的协作写书任务，可作为任务树原型。", None, None, 270000, teacher_id, now, now),
            )
            root_id = conn.execute("SELECT id FROM tasks WHERE title='太空能源技术概论（示例书稿）'").fetchone()[0]
            demo_tasks = [
                ("第1章 绪论", "约 1 万字。建立全书背景、定义与研究范畴。", root_id, s1, 10000, teacher_id, now, now),
                ("第2章 太空能源利用的核心挑战与关键问题", "约 0.5 万字。围绕来源、转换、传输三类挑战展开。", root_id, s2, 5000, teacher_id, now, now),
                ("第3章 太阳能直接发电技术", "约 2.5 万字。包含光伏与温差发电。", root_id, s1, 25000, teacher_id, now, now),
                ("第4章 空间热动力发电技术", "约 3 万字。", root_id, s2, 30000, teacher_id, now, now),
            ]
            conn.executemany(
                """
                INSERT INTO tasks(title, description, parent_id, assigned_to, target_words, created_by, created_at, updated_at)
                VALUES(?,?,?,?,?,?,?,?)
                """,
                demo_tasks,
            )

        admin_migration = conn.execute(
            "SELECT setting_value FROM system_settings WHERE setting_key='initial_admin_accounts_v1'"
        ).fetchone()
        if not admin_migration:
            now = datetime.now().isoformat(timespec="seconds")
            initial_admins = [
                ("Kuiruru", "Kuiruru", "Zxcv？6787"),
                ("AsterN7q", "AsterN7q", "M8!vR2#qL7@x"),
                ("VantaC4x", "VantaC4x", "Q6@pN9!sK3#t"),
            ]
            primary = conn.execute("SELECT * FROM users WHERE username='Kuiruru'").fetchone()
            legacy = conn.execute("SELECT * FROM users WHERE username='admin'").fetchone()
            if primary:
                conn.execute(
                    "UPDATE users SET display_name=?, role='admin', student_id=NULL, password_hash=?, updated_at=? WHERE id=?",
                    ("Kuiruru", password_hash("Zxcv？6787"), now, primary["id"]),
                )
                if legacy and legacy["id"] != primary["id"]:
                    legacy_username = f"legacy_{secrets.token_hex(5)}"
                    conn.execute(
                        "UPDATE users SET username=?, display_name=?, role='teacher', student_id=NULL, password_hash=?, updated_at=? WHERE id=?",
                        (legacy_username, "旧账号", password_hash(secrets.token_urlsafe(24)), now, legacy["id"]),
                    )
            elif legacy:
                conn.execute(
                    "UPDATE users SET username=?, display_name=?, role='admin', student_id=NULL, password_hash=?, updated_at=? WHERE id=?",
                    ("Kuiruru", "Kuiruru", password_hash("Zxcv？6787"), now, legacy["id"]),
                )
            else:
                conn.execute(
                    """
                    INSERT INTO users(username, display_name, role, password_hash, student_id, phone, email, created_at, updated_at)
                    VALUES(?,?,?,?,?,?,?,?,?)
                    """,
                    ("Kuiruru", "Kuiruru", "admin", password_hash("Zxcv？6787"), None, "", "", now, now),
                )

            for username, display_name, initial_password in initial_admins[1:]:
                existing = conn.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
                if existing:
                    conn.execute(
                        "UPDATE users SET display_name=?, role='admin', student_id=NULL, password_hash=?, updated_at=? WHERE id=?",
                        (display_name, password_hash(initial_password), now, existing["id"]),
                    )
                else:
                    conn.execute(
                        """
                        INSERT INTO users(username, display_name, role, password_hash, student_id, phone, email, created_at, updated_at)
                        VALUES(?,?,?,?,?,?,?,?,?)
                        """,
                        (username, display_name, "admin", password_hash(initial_password), None, "", "", now, now),
                    )
            conn.execute(
                "INSERT OR REPLACE INTO system_settings(setting_key, setting_value) VALUES('initial_admin_accounts_v1', ?)",
                (now,),
            )



def current_user(request: Request) -> sqlite3.Row | None:
    user_id = request.session.get("user_id")
    if not user_id:
        return None
    with db_conn() as conn:
        return conn.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()


def require_user(request: Request) -> sqlite3.Row:
    user = current_user(request)
    if not user:
        raise HTTPException(status_code=302, headers={"Location": "/login"})
    return user


def require_teacher_or_admin(user: sqlite3.Row) -> None:
    if user["role"] not in {"teacher", "admin"}:
        raise HTTPException(status_code=403, detail="没有权限执行该操作")


def require_admin(user: sqlite3.Row) -> None:
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="只有管理员账号可以执行该操作")


def valid_student_id(value: str) -> bool:
    value = (value or "").strip()
    return len(value) == 9 and value[2:3].isalpha() and value[2:3].isascii()


def next_teacher_username(conn: sqlite3.Connection) -> str:
    while True:
        username = f"teacher_{secrets.token_hex(3)}"
        exists = conn.execute("SELECT 1 FROM users WHERE username=?", (username,)).fetchone()
        if not exists:
            return username


def next_admin_username(conn: sqlite3.Connection) -> str:
    while True:
        username = f"N{secrets.token_hex(5)}"
        in_users = conn.execute("SELECT 1 FROM users WHERE username=?", (username,)).fetchone()
        in_requests = conn.execute("SELECT 1 FROM admin_creation_requests WHERE username=?", (username,)).fetchone()
        if not in_users and not in_requests:
            return username


def count_users_by_student_id(conn: sqlite3.Connection, student_id: str, exclude_user_id: int | None = None) -> int:
    if exclude_user_id is None:
        return conn.execute("SELECT COUNT(*) FROM users WHERE student_id=?", (student_id,)).fetchone()[0]
    return conn.execute("SELECT COUNT(*) FROM users WHERE student_id=? AND id<>?", (student_id, exclude_user_id)).fetchone()[0]


def count_words(text: str) -> int:
    """粗略统计中英文混合字数：中文按字，英文/数字按词。"""
    chinese_chars = re.findall(r"[\u4e00-\u9fff]", text or "")
    english_words = re.findall(r"[A-Za-z0-9]+(?:[-_'][A-Za-z0-9]+)*", text or "")
    return len(chinese_chars) + len(english_words)


def update_task_status(conn: sqlite3.Connection, task_id: int) -> None:
    task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
    if not task:
        return
    words = count_words(task["content"] or "")
    target = task["target_words"] or 0
    if target > 0:
        status = "已完成" if words >= target else "未完成"
    else:
        progress_info = assigned_task_block_progress(conn, task) if task["due_date"] else {"enabled": False, "percent": 0}
        status = "已完成" if progress_info.get("enabled") and progress_info.get("percent", 0) >= 100 else "未完成"
    conn.execute(
        "UPDATE tasks SET status=?, updated_at=? WHERE id=?",
        (status, datetime.now().isoformat(timespec="seconds"), task_id),
    )


def task_progress(task: sqlite3.Row) -> dict[str, Any]:
    words = count_words(task["content"] or "")
    target = task["target_words"] or 0
    percent = min(100, round(words / target * 100, 1)) if target else 0
    return {"words": words, "target": target, "percent": percent}


def tokenize(text: str) -> set[str]:
    text = text or ""
    chinese_chunks = set(re.findall(r"[\u4e00-\u9fff]{2,}", text))
    english = set(w.lower() for w in re.findall(r"[A-Za-z0-9]{3,}", text))
    return chinese_chunks | english


def similarity(a: str, b: str) -> float:
    ta, tb = tokenize(a), tokenize(b)
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta | tb)


def repeat_warnings(conn: sqlite3.Connection, task_id: int) -> list[dict[str, Any]]:
    task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
    if not task or not task["parent_id"]:
        return []
    siblings = conn.execute(
        "SELECT id, title, content FROM tasks WHERE parent_id=? AND id<>?",
        (task["parent_id"], task_id),
    ).fetchall()
    warnings: list[dict[str, Any]] = []
    for item in siblings:
        score = similarity(task["content"] or "", item["content"] or "")
        if score >= 0.18:
            warnings.append({"title": item["title"], "score": round(score * 100, 1)})
    return sorted(warnings, key=lambda x: x["score"], reverse=True)


def register_signin(user_id: int) -> None:
    today = date.today().isoformat()
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        conn.execute(
            "INSERT OR IGNORE INTO signins(user_id, signin_date, created_at) VALUES(?,?,?)",
            (user_id, today, now),
        )


def month_signin_count(user_id: int) -> int:
    prefix = date.today().strftime("%Y-%m")
    with db_conn() as conn:
        return conn.execute(
            "SELECT COUNT(*) FROM signins WHERE user_id=? AND signin_date LIKE ?",
            (user_id, f"{prefix}%"),
        ).fetchone()[0]


def month_signin_calendar(user_id: int) -> dict[str, Any]:
    today_obj = date.today()
    year, month = today_obj.year, today_obj.month
    days_in_month = calendar.monthrange(year, month)[1]
    first_weekday = date(year, month, 1).weekday()  # Monday=0
    prefix = f"{year:04d}-{month:02d}"
    with db_conn() as conn:
        signed = {
            row["signin_date"]
            for row in conn.execute(
                "SELECT signin_date FROM signins WHERE user_id=? AND signin_date LIKE ?",
                (user_id, f"{prefix}%"),
            ).fetchall()
        }
    days: list[dict[str, Any]] = []
    today_iso = today_obj.isoformat()
    for day in range(1, days_in_month + 1):
        current = date(year, month, day).isoformat()
        if current in signed:
            state = "signed"
        elif current > today_iso:
            state = "future"
        else:
            state = "missed"
        days.append({"date": current, "day": day, "state": state})
    return {
        "year": year,
        "month": month,
        "weekdays": ["一", "二", "三", "四", "五", "六", "日"],
        "leading_blanks": list(range(first_weekday)),
        "days": days,
        "signed_count": sum(1 for item in days if item["state"] == "signed"),
    }


def require_write_access(request: Request) -> None:
    if request.session.get("emergency_mode"):
        raise HTTPException(status_code=403, detail="紧急登录模式为只读状态。忘记密码的场合，请联系管理员！")


def validate_due_date_not_before_start(due_date: str, start_date: str) -> str | None:
    due = (due_date or "").strip()
    start = (start_date or "").strip()
    if due and start and due < start:
        return "截止日期不能早于开始日期"
    return None


def nav_context(request: Request) -> dict[str, Any]:
    user = current_user(request)
    return {
        "request": request,
        "user": user,
        "emergency_mode": bool(request.session.get("emergency_mode")),
        "show_emergency_alert": bool(request.session.pop("show_emergency_alert", False)),
    }


def safe_stored_name(suffix: str) -> str:
    return f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{secrets.token_hex(8)}{suffix}"


def ensure_docx_available() -> None:
    if DocxDocument is None:
        raise HTTPException(status_code=500, detail="缺少 python-docx 依赖。请重新运行 pip install -r requirements.txt")


def extract_text_from_docx(path: Path) -> str:
    ensure_docx_available()
    doc = DocxDocument(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def write_text_to_docx(path: Path, title: str, text: str) -> None:
    ensure_docx_available()
    doc = DocxDocument()
    doc.add_heading(title or "在线编辑文档", level=1)
    normalized = (text or "").replace("\r\n", "\n")
    for block in re.split(r"\n\s*\n", normalized):
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        if not lines:
            continue
        doc.add_paragraph("\n".join(lines))
    doc.save(str(path))


def extract_text_from_file(path: Path, suffix: str) -> str:
    suffix = suffix.lower()
    if suffix == ".docx":
        return extract_text_from_docx(path)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    return ""


def write_text_to_file(path: Path, suffix: str, title: str, text: str) -> None:
    suffix = suffix.lower()
    if suffix == ".docx":
        write_text_to_docx(path, title, text)
        return
    if suffix in {".txt", ".md"}:
        path.write_text(text or "", encoding="utf-8")
        return
    raise HTTPException(status_code=400, detail="当前在线保存仅支持 docx、txt、md 文件")


def document_word_count(item: sqlite3.Row) -> int:
    return count_words(item["extracted_text"] or "")


def can_edit_task(user: sqlite3.Row, task: sqlite3.Row | None) -> bool:
    if user["role"] in {"teacher", "admin"}:
        return True
    return bool(task and task["assigned_to"] == user["id"])


def get_document_for_user(conn: sqlite3.Connection, document_id: int, user: sqlite3.Row) -> sqlite3.Row:
    item = conn.execute(
        """
        SELECT d.*, t.title AS task_title, t.assigned_to AS task_assigned_to,
               u.display_name AS creator_name, uu.display_name AS updater_name
        FROM document_files d
        LEFT JOIN tasks t ON d.task_id=t.id
        LEFT JOIN users u ON d.created_by=u.id
        LEFT JOIN users uu ON d.updated_by=uu.id
        WHERE d.id=?
        """,
        (document_id,),
    ).fetchone()
    if not item:
        raise HTTPException(status_code=404, detail="文档不存在")
    # 当前测试阶段：老师/管理员可访问全部；学生可访问全部项目文档，但只能编辑分配给自己的任务文档或本人上传文档。
    return item


def can_edit_document(user: sqlite3.Row, item: sqlite3.Row) -> bool:
    if user["role"] in {"teacher", "admin"}:
        return True
    return item["created_by"] == user["id"] or item["task_assigned_to"] == user["id"]


def can_delete_document(user: sqlite3.Row, item: sqlite3.Row) -> bool:
    if user["role"] == "admin":
        return True
    if user["role"] == "student" and item["created_by"] == user["id"]:
        return True
    return False


def cleanup_file(path: Path) -> None:
    try:
        if path.exists():
            path.unlink()
    except OSError:
        pass


def delete_document_record(conn: sqlite3.Connection, document_id: int) -> None:
    item = conn.execute("SELECT * FROM document_files WHERE id=?", (document_id,)).fetchone()
    if not item:
        return
    cleanup_file(DOCUMENT_FILE_DIR / item["stored_name"])
    versions = conn.execute("SELECT stored_name FROM document_versions WHERE document_id=?", (document_id,)).fetchall()
    for version in versions:
        cleanup_file(DOCUMENT_VERSION_DIR / version["stored_name"])
    if item["task_id"]:
        task = conn.execute("SELECT content FROM tasks WHERE id=?", (item["task_id"],)).fetchone()
        if task and (task["content"] or "") == (item["extracted_text"] or ""):
            conn.execute("UPDATE tasks SET content='', updated_at=? WHERE id=?", (datetime.now().isoformat(timespec="seconds"), item["task_id"]))
            update_task_status(conn, item["task_id"])
    conn.execute("DELETE FROM document_versions WHERE document_id=?", (document_id,))
    conn.execute("DELETE FROM document_files WHERE id=?", (document_id,))


def task_accessible_for_images(user: sqlite3.Row, task: sqlite3.Row | None) -> bool:
    return can_edit_task(user, task)


def summarize_locally(text: str, max_sentences: int = 5) -> str:
    text = re.sub(r"\s+", " ", (text or "").strip())
    if not text:
        return "暂无可总结的正文内容。"
    parts = re.split(r"(?<=[。！？!?；;\.])\s*", text)
    sentences = [p.strip() for p in parts if len(p.strip()) >= 12]
    if not sentences:
        return text[:500] + ("……" if len(text) > 500 else "")
    tokens = tokenize(text)
    scored: list[tuple[int, int, str]] = []
    for idx, sentence in enumerate(sentences[:80]):
        stokens = tokenize(sentence)
        score = len(stokens & tokens) + min(len(sentence), 120) // 30
        scored.append((score, -idx, sentence))
    selected = sorted(scored, reverse=True)[:max_sentences]
    selected_sentences = [x[2] for x in sorted(selected, key=lambda x: -x[1])]
    summary = "\n".join(f"• {line}" for line in selected_sentences)
    return summary[:1200]


def summarize_with_ai_api(text: str) -> tuple[str, str]:
    """Optional OpenAI-compatible chat-completions style API. Falls back to local summary when not configured."""
    api_url = os.getenv("AI_API_URL", "").strip()
    api_key = os.getenv("AI_API_KEY", "").strip()
    model = os.getenv("AI_MODEL", "").strip() or "default"
    if not api_url or not api_key:
        return summarize_locally(text), "本地抽取式总结"
    prompt = (
        "请对下面的学术写作片段做简洁中文总结。要求：1. 用3-6条要点；"
        "2. 提炼核心论点、研究对象和可能与其他章节重复的主题；3. 不添加原文没有的信息。\n\n"
        f"正文：\n{text[:12000]}"
    )
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": "你是学术协作写作项目的文档总结助手。"},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.2,
    }
    req = urllib.request.Request(
        api_url,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        content = data.get("choices", [{}])[0].get("message", {}).get("content", "").strip()
        return content or summarize_locally(text), f"外部AI接口：{model}"
    except Exception as exc:
        return summarize_locally(text) + f"\n\n（外部AI接口暂不可用，已使用本地总结。错误：{type(exc).__name__}）", "本地抽取式总结"


def document_ai_warnings(conn: sqlite3.Connection, document_id: int) -> list[dict[str, Any]]:
    item = conn.execute(
        """
        SELECT d.*, t.parent_id, t.id AS linked_task_id, t.title AS task_title
        FROM document_files d LEFT JOIN tasks t ON d.task_id=t.id
        WHERE d.id=?
        """,
        (document_id,),
    ).fetchone()
    if not item or not item["parent_id"]:
        return []
    base_text = item["ai_summary"] or summarize_locally(item["extracted_text"] or "")
    rows = conn.execute(
        """
        SELECT d.id, d.title, d.ai_summary, d.extracted_text, t.title AS task_title, u.display_name AS creator_name
        FROM document_files d
        JOIN tasks t ON d.task_id=t.id
        LEFT JOIN users u ON d.created_by=u.id
        WHERE t.parent_id=? AND d.id<>? AND (d.deleted_at IS NULL OR d.deleted_at='')
        ORDER BY d.updated_at DESC
        """,
        (item["parent_id"], document_id),
    ).fetchall()
    warnings: list[dict[str, Any]] = []
    for row in rows:
        other_text = row["ai_summary"] or summarize_locally(row["extracted_text"] or "")
        score = similarity(base_text, other_text)
        if score >= 0.12:
            warnings.append({
                "document_id": row["id"],
                "title": row["title"],
                "task_title": row["task_title"],
                "creator_name": row["creator_name"],
                "score": round(score * 100, 1),
            })
    return sorted(warnings, key=lambda x: x["score"], reverse=True)


def require_student(user: sqlite3.Row) -> None:
    if user["role"] != "student":
        raise HTTPException(status_code=403, detail="该功能仅面向学生个人任务使用")


def personal_frequency_label(value: str) -> str:
    labels = {"daily": "日常", "weekly": "周常", "once": "一次性"}
    return labels.get(value, value or "未设置")


def personal_task_view(item: sqlite3.Row) -> dict[str, Any]:
    is_overdue = bool(item["due_date"] and item["status"] != "已完成" and item["due_date"] < date.today().isoformat())
    return {"task": item, "frequency_label": personal_frequency_label(item["frequency"]), "is_overdue": is_overdue}



def parse_iso_date(value: str | None) -> date | None:
    if not value:
        return None
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def week_bounds(today_obj: date) -> tuple[date, date]:
    start = today_obj - timedelta(days=today_obj.weekday())
    return start, start + timedelta(days=6)


def generate_recurring_personal_tasks(user_id: int, conn: sqlite3.Connection | None = None) -> None:
    """为启用自动续订的日常/周常任务生成当前周期的个人任务实例。"""
    close_after = False
    if conn is None:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        close_after = True
    try:
        today_obj = date.today()
        today_iso = today_obj.isoformat()
        now = datetime.now().isoformat(timespec="seconds")
        conn.execute(
            """
            UPDATE personal_tasks
            SET status='已完成', completed_at=COALESCE(completed_at, ?), updated_at=?
            WHERE owner_id=? AND COALESCE(is_template,0)=0 AND frequency='daily'
              AND status!='已完成' AND due_date IS NOT NULL AND due_date < ?
            """,
            (now, now, user_id, today_iso),
        )
        templates_rows = conn.execute(
            """
            SELECT * FROM personal_tasks
            WHERE owner_id=? AND is_template=1 AND auto_renew=1 AND frequency IN ('daily','weekly')
            ORDER BY id
            """,
            (user_id,),
        ).fetchall()
        for tmpl in templates_rows:
            start_obj = parse_iso_date(tmpl["start_date"] or tmpl["created_at"])
            end_obj = parse_iso_date(tmpl["due_date"])
            if start_obj and today_obj < start_obj:
                continue
            if end_obj and today_obj > end_obj:
                continue
            if tmpl["frequency"] == "daily":
                period_start = today_obj
                period_end = today_obj
                key = today_iso
            else:
                raw_start, raw_end = week_bounds(today_obj)
                period_start = max(raw_start, start_obj or raw_start)
                period_end = min(raw_end, end_obj or raw_end)
                iso_year, iso_week, _ = today_obj.isocalendar()
                key = f"{iso_year}-W{iso_week:02d}"
            exists = conn.execute(
                "SELECT 1 FROM personal_tasks WHERE recurrence_parent_id=? AND occurrence_key=?",
                (tmpl["id"], key),
            ).fetchone()
            if exists:
                continue
            conn.execute(
                """
                INSERT INTO personal_tasks(owner_id, title, description, frequency, start_date, status, due_date,
                                           created_at, updated_at, is_template, auto_renew, recurrence_parent_id, occurrence_key)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    user_id,
                    tmpl["title"],
                    tmpl["description"] or "",
                    tmpl["frequency"],
                    period_start.isoformat(),
                    "未完成",
                    period_end.isoformat(),
                    now,
                    now,
                    0,
                    0,
                    tmpl["id"],
                    key,
                ),
            )
        if close_after:
            conn.commit()
    finally:
        if close_after:
            conn.close()


def build_time_blocks(start_value: str | None, due_value: str | None) -> list[dict[str, Any]]:
    start = parse_iso_date(start_value)
    due = parse_iso_date(due_value)
    if not start or not due or due < start:
        return []
    total_days = (due - start).days + 1
    if total_days <= 7:
        return [{"index": 0, "start": start.isoformat(), "end": due.isoformat(), "label": f"{start.isoformat()} ～ {due.isoformat()}"}]
    blocks: list[dict[str, Any]] = []
    current = start
    idx = 0
    while current <= due:
        block_end = min(current + timedelta(days=6), due)
        blocks.append({"index": idx, "start": current.isoformat(), "end": block_end.isoformat(), "label": f"第 {idx + 1} 周：{current.isoformat()} ～ {block_end.isoformat()}"})
        current = block_end + timedelta(days=1)
        idx += 1
    return blocks


def current_block_index(blocks: list[dict[str, Any]]) -> int:
    if not blocks:
        return 0
    today_iso = date.today().isoformat()
    for block in blocks:
        if block["start"] <= today_iso <= block["end"]:
            return int(block["index"])
    if today_iso < blocks[0]["start"]:
        return 0
    return int(blocks[-1]["index"])


def fetch_progress_records(conn: sqlite3.Connection, record_type: str, task_id: int) -> list[sqlite3.Row]:
    return conn.execute(
        """
        SELECT r.*, u.display_name AS owner_name
        FROM task_progress_records r LEFT JOIN users u ON r.owner_id=u.id
        WHERE r.record_type=? AND r.task_id=?
        ORDER BY r.block_index, r.created_at DESC, r.id DESC
        """,
        (record_type, task_id),
    ).fetchall()


def block_progress_from_records(blocks: list[dict[str, Any]], records: list[sqlite3.Row]) -> dict[str, Any]:
    record_map: dict[int, list[sqlite3.Row]] = {}
    for record in records:
        record_map.setdefault(int(record["block_index"]), []).append(record)
    for block in blocks:
        block["records"] = record_map.get(int(block["index"]), [])
        block["recorded"] = bool(block["records"])
        block["is_current"] = int(block["index"]) == current_block_index(blocks)
    if not blocks:
        return {"enabled": False, "blocks": [], "percent": 0, "record_count": len(records), "current_index": 0}
    valid_indexes = [k for k in record_map.keys() if 0 <= k < len(blocks)]
    if valid_indexes:
        max_index = max(valid_indexes)
        percent = round((max_index + 1) / len(blocks) * 100, 1)
    else:
        percent = 0
    return {"enabled": True, "blocks": blocks, "percent": percent, "record_count": len(records), "current_index": current_block_index(blocks)}


def milestone_progress_from_records(blocks: list[dict[str, Any]], records: list[sqlite3.Row]) -> dict[str, Any]:
    """一次性个人任务使用的里程碑进度。

    未来节点只展示，不开放填写；只有到达每个 7 天节点或最终截止日后，
    才开放对应的“本周进展”窗口。保存文字记录后，进度按该节点累计天数 / 总天数推进。
    """
    record_map: dict[int, list[sqlite3.Row]] = {}
    for record in records:
        record_map.setdefault(int(record["block_index"]), []).append(record)
    if not blocks:
        return {
            "enabled": False,
            "blocks": [],
            "visible_blocks": [],
            "percent": 0,
            "record_count": len(records),
            "current_index": None,
            "current_block": None,
            "next_block": None,
            "total_days": 0,
            "progress_days": 0,
        }

    start = parse_iso_date(blocks[0]["start"])
    due = parse_iso_date(blocks[-1]["end"])
    today = date.today()
    total_days = ((due - start).days + 1) if start and due else len(blocks) * 7
    current_index: int | None = None
    next_block: dict[str, Any] | None = None

    for idx, block in enumerate(blocks):
        block_start = parse_iso_date(block["start"])
        block_end = parse_iso_date(block["end"])
        next_unlock = parse_iso_date(blocks[idx + 1]["end"]) if idx + 1 < len(blocks) else None
        unlock_date = block_end
        report_until = (next_unlock - timedelta(days=1)) if next_unlock else block_end
        duration_days = ((block_end - block_start).days + 1) if block_start and block_end else 0
        cumulative_days = ((block_end - start).days + 1) if start and block_end else 0

        block["records"] = record_map.get(int(block["index"]), [])
        block["recorded"] = bool(block["records"])
        block["unlock_date"] = unlock_date.isoformat() if unlock_date else ""
        block["report_until"] = report_until.isoformat() if report_until else ""
        block["duration_days"] = duration_days
        block["cumulative_days"] = min(cumulative_days, total_days)
        block["percent_after"] = round(min(cumulative_days, total_days) / total_days * 100, 2) if total_days else 0
        block["is_future"] = bool(unlock_date and today < unlock_date)
        block["is_open"] = bool(unlock_date and report_until and unlock_date <= today <= report_until)
        block["is_expired"] = bool(report_until and today > report_until and not block["recorded"])
        block["is_current"] = block["is_open"]
        if block["is_open"]:
            current_index = int(block["index"])
        elif unlock_date and today < unlock_date and next_block is None:
            next_block = block

    valid_indexes = [k for k in record_map.keys() if 0 <= k < len(blocks)]
    if valid_indexes:
        max_index = max(valid_indexes)
        progress_days = int(blocks[max_index].get("cumulative_days") or 0)
        percent = round(progress_days / total_days * 100, 2) if total_days else 0
    else:
        progress_days = 0
        percent = 0

    current_block = blocks[current_index] if current_index is not None else None
    # 只显示已到达节点、已记录节点和下一次即将开放的节点，避免未来所有区块铺满页面。
    visible_blocks = [b for b in blocks if b["recorded"] or not b["is_future"]]
    if next_block and next_block not in visible_blocks:
        visible_blocks.append(next_block)

    return {
        "enabled": True,
        "blocks": blocks,
        "visible_blocks": visible_blocks,
        "percent": percent,
        "record_count": len(records),
        "current_index": current_index,
        "current_block": current_block,
        "next_block": next_block,
        "total_days": total_days,
        "progress_days": progress_days,
    }


def assigned_task_block_progress(conn: sqlite3.Connection, task: sqlite3.Row) -> dict[str, Any]:
    if (task["target_words"] or 0) > 0 or not task["due_date"]:
        return {"enabled": False, "blocks": [], "percent": 0, "record_count": 0, "current_index": 0}
    blocks = build_time_blocks((task["created_at"] or date.today().isoformat())[:10], task["due_date"])
    records = fetch_progress_records(conn, "assigned", task["id"])
    return block_progress_from_records(blocks, records)


def personal_once_block_progress(conn: sqlite3.Connection, item: sqlite3.Row) -> dict[str, Any]:
    if item["frequency"] != "once" or item["is_template"]:
        return {
            "enabled": False,
            "blocks": [],
            "visible_blocks": [],
            "percent": 0,
            "record_count": 0,
            "current_index": None,
            "current_block": None,
            "next_block": None,
            "total_days": 0,
            "progress_days": 0,
        }
    blocks = build_time_blocks(item["start_date"] or (item["created_at"] or "")[:10], item["due_date"])
    records = fetch_progress_records(conn, "personal", item["id"])
    return milestone_progress_from_records(blocks, records)


def task_progress_for_card(conn: sqlite3.Connection, task: sqlite3.Row) -> dict[str, Any]:
    base = task_progress(task)
    if (task["target_words"] or 0) <= 0:
        block_progress = assigned_task_block_progress(conn, task)
        if block_progress["enabled"]:
            return {"words": 0, "target": 0, "percent": block_progress["percent"], "mode": "block", "blocks_total": len(block_progress["blocks"]), "record_count": block_progress["record_count"]}
    base["mode"] = "words"
    return base


def update_block_based_task_status(conn: sqlite3.Connection, task_id: int) -> None:
    task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
    if not task or (task["target_words"] or 0) > 0:
        return
    progress_info = assigned_task_block_progress(conn, task)
    if progress_info["enabled"]:
        status = "已完成" if progress_info["percent"] >= 100 else "未完成"
        conn.execute("UPDATE tasks SET status=?, updated_at=? WHERE id=?", (status, datetime.now().isoformat(timespec="seconds"), task_id))


def save_progress_upload(file: UploadFile | None) -> tuple[str, str, str]:
    if file is None or not file.filename:
        return "", "", ""
    suffix = Path(file.filename).suffix.lower()
    allowed = {".docx", ".txt", ".md", ".pdf", ".pptx", ".xlsx", ".csv", ".png", ".jpg", ".jpeg", ".webp", ".gif"}
    if suffix not in allowed:
        raise HTTPException(status_code=400, detail="进度记录附件仅支持常见文档、表格、图片格式")
    stored_name = safe_stored_name(suffix)
    dest = PROJECT_FILE_DIR / stored_name
    with dest.open("wb") as f:
        shutil.copyfileobj(file.file, f)
    mime_type = file.content_type or mimetypes.guess_type(file.filename)[0] or "application/octet-stream"
    return file.filename, stored_name, mime_type


def progress_record_access(conn: sqlite3.Connection, record: sqlite3.Row, user: sqlite3.Row) -> bool:
    if user["role"] in {"teacher", "admin"}:
        return True
    if record["record_type"] == "personal":
        return record["owner_id"] == user["id"]
    task = conn.execute("SELECT assigned_to FROM tasks WHERE id=?", (record["task_id"],)).fetchone()
    return bool(task and task["assigned_to"] == user["id"])


def add_auto_assigned_progress_record(conn: sqlite3.Connection, task_id: int, user_id: int, note: str) -> None:
    task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
    if not task or (task["target_words"] or 0) > 0:
        return
    progress_info = assigned_task_block_progress(conn, task)
    if not progress_info["enabled"]:
        return
    idx = progress_info["current_index"]
    blocks = progress_info["blocks"]
    if idx < 0 or idx >= len(blocks):
        return
    block = blocks[idx]
    now = datetime.now().isoformat(timespec="seconds")
    conn.execute(
        """
        INSERT INTO task_progress_records(record_type, task_id, owner_id, block_index, block_start, block_end, note, created_at)
        VALUES(?,?,?,?,?,?,?,?)
        """,
        ("assigned", task_id, user_id, idx, block["start"], block["end"], note, now),
    )
    update_task_status(conn, task_id)


@app.on_event("startup")
def startup() -> None:
    init_db()


@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    user = current_user(request)
    if not user:
        return RedirectResponse("/login", status_code=303)

    with db_conn() as conn:
        if user["role"] in {"teacher", "admin"}:
            tasks = conn.execute(
                """
                SELECT t.*, u.display_name AS assigned_name
                FROM tasks t LEFT JOIN users u ON t.assigned_to=u.id
                ORDER BY COALESCE(t.parent_id, t.id), t.parent_id IS NOT NULL, t.id
                """
            ).fetchall()
        else:
            tasks = conn.execute(
                """
                SELECT t.*, u.display_name AS assigned_name
                FROM tasks t LEFT JOIN users u ON t.assigned_to=u.id
                WHERE t.assigned_to=?
                ORDER BY t.due_date IS NULL, t.due_date, t.id
                """,
                (user["id"],),
            ).fetchall()
        task_cards = [{"task": t, "progress": task_progress_for_card(conn, t)} for t in tasks]
        reminders = [x for x in task_cards if x["task"]["status"] != "已完成"]
        total_tasks = len(task_cards)
        done_tasks = sum(1 for x in task_cards if x["task"]["status"] == "已完成")
        avg_progress = round(sum(x["progress"]["percent"] for x in task_cards) / total_tasks, 1) if total_tasks else 0
        today = date.today().isoformat()
        overdue_tasks = sum(
            1
            for x in task_cards
            if x["task"]["status"] != "已完成" and x["task"]["due_date"] and x["task"]["due_date"] < today
        )
        document_count = conn.execute("SELECT COUNT(*) FROM document_files").fetchone()[0]
        if user["role"] == "student":
            generate_recurring_personal_tasks(user["id"], conn)
            personal_rows = conn.execute(
                """
                SELECT * FROM personal_tasks
                WHERE owner_id=? AND status!='已完成' AND COALESCE(is_template,0)=0
                ORDER BY due_date IS NULL, due_date, frequency, id
                """,
                (user["id"],),
            ).fetchall()
            if request.session.pop("show_overdue_personal_check", False) and not request.session.get("emergency_mode"):
                overdue_popup_rows = conn.execute(
                    """
                    SELECT * FROM personal_tasks
                    WHERE owner_id=? AND status!='已完成' AND COALESCE(is_template,0)=0
                      AND frequency!='daily' AND due_date IS NOT NULL AND due_date < ?
                    ORDER BY due_date, id
                    """,
                    (user["id"], today),
                ).fetchall()
            else:
                overdue_popup_rows = []
        else:
            personal_rows = []
            overdue_popup_rows = []
        personal_reminders = [personal_task_view(x) for x in personal_rows]
        overdue_popup_tasks = [personal_task_view(x) for x in overdue_popup_rows]
        stats = {
            "total_tasks": total_tasks,
            "done_tasks": done_tasks,
            "active_tasks": total_tasks - done_tasks,
            "avg_progress": avg_progress,
            "overdue_tasks": overdue_tasks,
            "document_count": document_count,
        }

    ctx = nav_context(request)
    ctx.update(
        {
            "task_cards": task_cards,
            "reminders": reminders,
            "personal_reminders": personal_reminders,
            "signin_count": month_signin_count(user["id"]),
            "signin_calendar": month_signin_calendar(user["id"]),
            "overdue_popup_tasks": overdue_popup_tasks,
            "stats": stats,
        }
    )
    return templates.TemplateResponse(request, "index.html", ctx)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, registered: int = 0):
    message = "学生账号创建成功，请登录。" if registered else None
    return templates.TemplateResponse(request, "login.html", {"error": None, "message": message})


@app.post("/login")
def login(request: Request, username: str = Form(...), password: str = Form(...)):
    with db_conn() as conn:
        login_id = username.strip()
        user = conn.execute(
            "SELECT * FROM users WHERE username=? OR student_id=?",
            (login_id, login_id),
        ).fetchone()
    if not user or not verify_password(password, user["password_hash"]):
        return templates.TemplateResponse(request, "login.html", {"error": "用户名或密码错误"}, status_code=401)
    request.session.clear()
    request.session["user_id"] = user["id"]
    request.session["emergency_mode"] = False
    if user["role"] == "student":
        with db_conn() as conn:
            generate_recurring_personal_tasks(user["id"], conn)
        request.session["show_overdue_personal_check"] = True
    register_signin(user["id"])
    return RedirectResponse("/", status_code=303)


@app.post("/emergency-login")
def emergency_login(request: Request, username: str = Form(...)):
    login_id = username.strip()
    with db_conn() as conn:
        user = conn.execute(
            "SELECT * FROM users WHERE username=? OR student_id=?",
            (login_id, login_id),
        ).fetchone()
    if not user:
        return templates.TemplateResponse(request, "login.html", {"error": "未找到该用户名", "message": None}, status_code=401)
    request.session.clear()
    request.session["user_id"] = user["id"]
    request.session["emergency_mode"] = True
    request.session["show_emergency_alert"] = True
    if user["role"] == "student":
        with db_conn() as conn:
            generate_recurring_personal_tasks(user["id"], conn)
    register_signin(user["id"])
    return RedirectResponse("/", status_code=303)


@app.get("/register", response_class=HTMLResponse)
def register_student_page(request: Request):
    if current_user(request):
        return RedirectResponse("/", status_code=303)
    return templates.TemplateResponse(request, "register_student.html", {"request": request, "user": None, "error": None, "values": {}})


@app.post("/register")
def register_student(
    request: Request,
    display_name: str = Form(...),
    student_id: str = Form(...),
    password: str = Form(...),
    password_confirm: str = Form(...),
    phone: str = Form(""),
    email: str = Form(""),
):
    if current_user(request):
        return RedirectResponse("/", status_code=303)
    values = {
        "display_name": display_name.strip(),
        "student_id": student_id.strip(),
        "phone": phone.strip(),
        "email": email.strip(),
    }
    if not values["display_name"]:
        return templates.TemplateResponse(request, "register_student.html", {"request": request, "user": None, "error": "姓名不能为空", "values": values}, status_code=400)
    if not values["student_id"] or not valid_student_id(values["student_id"]):
        return templates.TemplateResponse(request, "register_student.html", {"request": request, "user": None, "error": "学号格式不正确", "values": values}, status_code=400)
    if not password:
        return templates.TemplateResponse(request, "register_student.html", {"request": request, "user": None, "error": "密码不能为空", "values": values}, status_code=400)
    if password != password_confirm:
        return templates.TemplateResponse(request, "register_student.html", {"request": request, "user": None, "error": "两次输入的密码不一致", "values": values}, status_code=400)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        if count_users_by_student_id(conn, values["student_id"]) > 0 or conn.execute("SELECT 1 FROM users WHERE username=?", (values["student_id"],)).fetchone():
            return templates.TemplateResponse(request, "register_student.html", {"request": request, "user": None, "error": "该学号已被使用，请使用其他学号", "values": values}, status_code=400)
        conn.execute(
            """
            INSERT INTO users(username, display_name, role, password_hash, student_id, phone, email, created_at, updated_at)
            VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (values["student_id"], values["display_name"], "student", password_hash(password), values["student_id"], values["phone"], values["email"], now, now),
        )
    return RedirectResponse("/login?registered=1", status_code=303)


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.get("/tasks/new", response_class=HTMLResponse)
def new_task_page(request: Request, parent_id: int | None = None):
    user = require_user(request)
    require_teacher_or_admin(user)
    with db_conn() as conn:
        students = conn.execute("SELECT id, display_name FROM users WHERE role='student' ORDER BY display_name").fetchall()
        parent_tasks = conn.execute("SELECT id, title FROM tasks ORDER BY id").fetchall()
    ctx = nav_context(request)
    ctx.update({"students": students, "parent_tasks": parent_tasks, "parent_id": parent_id})
    return templates.TemplateResponse(request, "task_form.html", ctx)


@app.post("/tasks/new")
def create_task(
    request: Request,
    title: str = Form(...),
    description: str = Form(""),
    parent_id: str = Form(""),
    assigned_to: str = Form(""),
    target_words: int = Form(0),
    due_date: str = Form(""),
):
    user = require_user(request)
    require_write_access(request)
    require_teacher_or_admin(user)
    if not due_date:
        raise HTTPException(status_code=400, detail="请为任务指定截止日期")
    parent = int(parent_id) if parent_id else None
    assignee = int(assigned_to) if assigned_to else None
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        conn.execute(
            """
            INSERT INTO tasks(title, description, parent_id, assigned_to, target_words, due_date, created_by, created_at, updated_at)
            VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (title.strip(), description.strip(), parent, assignee, target_words, due_date or None, user["id"], now, now),
        )
    return RedirectResponse("/", status_code=303)


@app.get("/tasks/{task_id}", response_class=HTMLResponse)
def task_detail(request: Request, task_id: int):
    user = require_user(request)
    with db_conn() as conn:
        task = conn.execute(
            """
            SELECT t.*, u.display_name AS assigned_name
            FROM tasks t LEFT JOIN users u ON t.assigned_to=u.id
            WHERE t.id=?
            """,
            (task_id,),
        ).fetchone()
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        if user["role"] == "student" and task["assigned_to"] != user["id"]:
            raise HTTPException(status_code=403, detail="你只能查看分配给自己的任务")
        children = conn.execute(
            """
            SELECT t.*, u.display_name AS assigned_name
            FROM tasks t LEFT JOIN users u ON t.assigned_to=u.id
            WHERE t.parent_id=? ORDER BY t.id
            """,
            (task_id,),
        ).fetchall()
        task_documents = conn.execute(
            """
            SELECT d.*, u.display_name AS creator_name
            FROM document_files d LEFT JOIN users u ON d.created_by=u.id
            WHERE d.task_id=? AND (d.deleted_at IS NULL OR d.deleted_at='') ORDER BY d.updated_at DESC
            """,
            (task_id,),
        ).fetchall()
        task_images = conn.execute(
            """
            SELECT i.*, u.display_name AS uploader_name
            FROM task_images i LEFT JOIN users u ON i.uploaded_by=u.id
            WHERE i.task_id=? ORDER BY i.uploaded_at DESC, i.id DESC
            """,
            (task_id,),
        ).fetchall()
        warnings = repeat_warnings(conn, task_id)
        block_progress = assigned_task_block_progress(conn, task)
        child_cards = [{"task": child, "progress": task_progress_for_card(conn, child)} for child in children]
        task_progress_card = task_progress_for_card(conn, task)
    doc_cards = [{"doc": doc, "word_count": document_word_count(doc)} for doc in task_documents]
    image_upload_allowed = task_accessible_for_images(user, task) and not request.session.get("emergency_mode")
    ctx = nav_context(request)
    ctx.update({
        "task": task,
        "children": children,
        "child_cards": child_cards,
        "progress": task_progress_card,
        "block_progress": block_progress,
        "warnings": warnings,
        "doc_cards": doc_cards,
        "task_images": task_images,
        "image_upload_allowed": image_upload_allowed,
    })
    return templates.TemplateResponse(request, "task_detail.html", ctx)


@app.get("/tasks/{task_id}/edit", response_class=HTMLResponse)
def edit_task_page(request: Request, task_id: int):
    user = require_user(request)
    with db_conn() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        if user["role"] == "student" and task["assigned_to"] != user["id"]:
            raise HTTPException(status_code=403, detail="你只能编辑分配给自己的任务")
        progress_card = task_progress_for_card(conn, task)
    ctx = nav_context(request)
    ctx.update({"task": task, "progress": progress_card, "warnings": []})
    return templates.TemplateResponse(request, "task_edit.html", ctx)


@app.post("/tasks/{task_id}/edit")
def edit_task(request: Request, task_id: int, content: str = Form(...)):
    user = require_user(request)
    require_write_access(request)
    with db_conn() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        if user["role"] == "student" and task["assigned_to"] != user["id"]:
            raise HTTPException(status_code=403, detail="你只能编辑分配给自己的任务")
        conn.execute("UPDATE tasks SET content=?, updated_at=? WHERE id=?", (content, datetime.now().isoformat(timespec="seconds"), task_id))
        update_task_status(conn, task_id)
    return RedirectResponse(f"/tasks/{task_id}", status_code=303)


@app.get("/documents", response_class=HTMLResponse)
def document_list(request: Request):
    user = require_user(request)
    with db_conn() as conn:
        docs = conn.execute(
            """
            SELECT d.*, t.title AS task_title, t.assigned_to AS task_assigned_to,
                   u.display_name AS creator_name, uu.display_name AS updater_name
            FROM document_files d
            LEFT JOIN tasks t ON d.task_id=t.id
            LEFT JOIN users u ON d.created_by=u.id
            LEFT JOIN users uu ON d.updated_by=uu.id
            WHERE d.deleted_at IS NULL OR d.deleted_at=''
            ORDER BY d.updated_at DESC, d.id DESC
            """
        ).fetchall()
        if user["role"] in {"teacher", "admin"}:
            available_tasks = conn.execute("SELECT id, title FROM tasks ORDER BY id").fetchall()
        else:
            available_tasks = conn.execute("SELECT id, title FROM tasks WHERE assigned_to=? ORDER BY id", (user["id"],)).fetchall()
    doc_cards = [{"doc": doc, "word_count": document_word_count(doc), "editable": can_edit_document(user, doc), "deletable": can_delete_document(user, doc)} for doc in docs]
    ctx = nav_context(request)
    ctx.update({"doc_cards": doc_cards, "available_tasks": available_tasks})
    return templates.TemplateResponse(request, "documents.html", ctx)


@app.post("/documents/upload")
def upload_document(
    request: Request,
    title: str = Form(""),
    description: str = Form(""),
    task_id: str = Form(""),
    file: UploadFile = File(...),
):
    user = require_user(request)
    require_write_access(request)
    suffix = Path(file.filename or "").suffix.lower()
    allowed = {".docx", ".txt", ".md"}
    if suffix not in allowed:
        raise HTTPException(status_code=400, detail="当前在线编辑仅支持 docx、txt、md。pptx/pdf 可继续放在模板库下载。")
    linked_task_id = int(task_id) if task_id else None
    with db_conn() as conn:
        linked_task = conn.execute("SELECT * FROM tasks WHERE id=?", (linked_task_id,)).fetchone() if linked_task_id else None
        if linked_task_id and not linked_task:
            raise HTTPException(status_code=404, detail="绑定任务不存在")
        if linked_task and not can_edit_task(user, linked_task):
            raise HTTPException(status_code=403, detail="你没有权限给该任务上传文档")

    stored_name = safe_stored_name(suffix)
    dest = DOCUMENT_FILE_DIR / stored_name
    with dest.open("wb") as f:
        shutil.copyfileobj(file.file, f)
    extracted = extract_text_from_file(dest, suffix)
    doc_title = title.strip() or Path(file.filename or "未命名文档").stem
    now = datetime.now().isoformat(timespec="seconds")

    with db_conn() as conn:
        conn.execute(
            """
            INSERT INTO document_files(title, original_name, stored_name, file_type, task_id, extracted_text,
                                       created_by, updated_by, uploaded_at, updated_at, description, current_version)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (doc_title, file.filename, stored_name, suffix.lstrip("."), linked_task_id, extracted, user["id"], user["id"], now, now, description.strip(), 1),
        )
        document_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        version_name = f"doc{document_id}_v1_{stored_name}"
        shutil.copy2(dest, DOCUMENT_VERSION_DIR / version_name)
        conn.execute(
            """
            INSERT INTO document_versions(document_id, version_no, stored_name, saved_by, saved_at, word_count, note)
            VALUES(?,?,?,?,?,?,?)
            """,
            (document_id, 1, version_name, user["id"], now, count_words(extracted), "初始上传"),
        )
        if linked_task_id:
            conn.execute("UPDATE tasks SET content=?, updated_at=? WHERE id=?", (extracted, now, linked_task_id))
            update_task_status(conn, linked_task_id)
            add_auto_assigned_progress_record(conn, linked_task_id, user["id"], f"上传并绑定在线文档：{doc_title}")
    return RedirectResponse(f"/documents/{document_id}/edit", status_code=303)


@app.get("/documents/{document_id}/edit", response_class=HTMLResponse)
def edit_document_page(request: Request, document_id: int):
    user = require_user(request)
    with db_conn() as conn:
        item = get_document_for_user(conn, document_id, user)
        versions = conn.execute(
            """
            SELECT v.*, u.display_name AS saver_name
            FROM document_versions v LEFT JOIN users u ON v.saved_by=u.id
            WHERE v.document_id=? ORDER BY v.version_no DESC
            """,
            (document_id,),
        ).fetchall()
        warnings = repeat_warnings(conn, item["task_id"]) if item["task_id"] else []
        ai_warnings = document_ai_warnings(conn, document_id)
    ctx = nav_context(request)
    ctx.update({
        "doc": item,
        "versions": versions,
        "word_count": document_word_count(item),
        "editable": can_edit_document(user, item) and not request.session.get("emergency_mode"),
        "warnings": warnings,
        "ai_warnings": ai_warnings,
    })
    return templates.TemplateResponse(request, "document_edit.html", ctx)


@app.post("/documents/{document_id}/edit")
def save_document(request: Request, document_id: int, content: str = Form(...), note: str = Form("")):
    user = require_user(request)
    require_write_access(request)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        item = get_document_for_user(conn, document_id, user)
        if not can_edit_document(user, item):
            raise HTTPException(status_code=403, detail="你没有权限编辑该文档")
        path = DOCUMENT_FILE_DIR / item["stored_name"]
        if not path.exists():
            raise HTTPException(status_code=404, detail="文档文件丢失")

        next_version = (item["current_version"] or 1) + 1
        suffix = "." + item["file_type"].lower().lstrip(".")
        write_text_to_file(path, suffix, item["title"], content)
        version_name = f"doc{document_id}_v{next_version}_{safe_stored_name(suffix)}"
        shutil.copy2(path, DOCUMENT_VERSION_DIR / version_name)
        words = count_words(content)
        conn.execute(
            """
            UPDATE document_files
            SET extracted_text=?, updated_by=?, updated_at=?, current_version=?
            WHERE id=?
            """,
            (content, user["id"], now, next_version, document_id),
        )
        conn.execute(
            """
            INSERT INTO document_versions(document_id, version_no, stored_name, saved_by, saved_at, word_count, note)
            VALUES(?,?,?,?,?,?,?)
            """,
            (document_id, next_version, version_name, user["id"], now, words, note.strip() or "在线编辑保存"),
        )
        if item["task_id"]:
            conn.execute("UPDATE tasks SET content=?, updated_at=? WHERE id=?", (content, now, item["task_id"]))
            update_task_status(conn, item["task_id"])
            add_auto_assigned_progress_record(conn, item["task_id"], user["id"], f"编辑并保存在线文档：{item['title']}")
    return RedirectResponse(f"/documents/{document_id}/edit", status_code=303)


@app.get("/documents/{document_id}/download")
def download_document(request: Request, document_id: int):
    user = require_user(request)
    with db_conn() as conn:
        item = get_document_for_user(conn, document_id, user)
    path = DOCUMENT_FILE_DIR / item["stored_name"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="文档文件丢失")
    suffix = "." + item["file_type"].lower().lstrip(".")
    filename = f"{item['title']}{suffix}"
    return FileResponse(path, filename=filename, media_type="application/octet-stream")


@app.get("/documents/{document_id}/versions/{version_id}/download")
def download_document_version(request: Request, document_id: int, version_id: int):
    user = require_user(request)
    with db_conn() as conn:
        item = get_document_for_user(conn, document_id, user)
        version = conn.execute(
            "SELECT * FROM document_versions WHERE id=? AND document_id=?",
            (version_id, document_id),
        ).fetchone()
    if not version:
        raise HTTPException(status_code=404, detail="历史版本不存在")
    path = DOCUMENT_VERSION_DIR / version["stored_name"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="历史版本文件丢失")
    suffix = "." + item["file_type"].lower().lstrip(".")
    filename = f"{item['title']}_v{version['version_no']}{suffix}"
    return FileResponse(path, filename=filename, media_type="application/octet-stream")


@app.post("/documents/{document_id}/delete")
def delete_document(request: Request, document_id: int):
    user = require_user(request)
    require_write_access(request)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        item = get_document_for_user(conn, document_id, user)
        if not can_delete_document(user, item):
            raise HTTPException(status_code=403, detail="只有文档上传者本人或管理员可以删除该文档")
        conn.execute("UPDATE document_files SET deleted_at=?, deleted_by=?, updated_at=? WHERE id=?", (now, user["id"], now, document_id))
        if item["task_id"]:
            task = conn.execute("SELECT content FROM tasks WHERE id=?", (item["task_id"],)).fetchone()
            if task and (task["content"] or "") == (item["extracted_text"] or ""):
                conn.execute("UPDATE tasks SET content='', updated_at=? WHERE id=?", (now, item["task_id"]))
                update_task_status(conn, item["task_id"])
    return RedirectResponse("/documents", status_code=303)


@app.get("/documents/trash", response_class=HTMLResponse)
def document_trash(request: Request):
    user = require_user(request)
    with db_conn() as conn:
        if user["role"] in {"teacher", "admin"}:
            docs = conn.execute(
                """
                SELECT d.*, t.title AS task_title, t.assigned_to AS task_assigned_to,
                       u.display_name AS creator_name, du.display_name AS deleter_name
                FROM document_files d
                LEFT JOIN tasks t ON d.task_id=t.id
                LEFT JOIN users u ON d.created_by=u.id
                LEFT JOIN users du ON d.deleted_by=du.id
                WHERE d.deleted_at IS NOT NULL AND d.deleted_at<>''
                ORDER BY d.deleted_at DESC, d.id DESC
                """
            ).fetchall()
        else:
            docs = conn.execute(
                """
                SELECT d.*, t.title AS task_title, t.assigned_to AS task_assigned_to,
                       u.display_name AS creator_name, du.display_name AS deleter_name
                FROM document_files d
                LEFT JOIN tasks t ON d.task_id=t.id
                LEFT JOIN users u ON d.created_by=u.id
                LEFT JOIN users du ON d.deleted_by=du.id
                WHERE d.created_by=? AND d.deleted_at IS NOT NULL AND d.deleted_at<>''
                ORDER BY d.deleted_at DESC, d.id DESC
                """,
                (user["id"],),
            ).fetchall()
    doc_cards = [{"doc": doc, "word_count": document_word_count(doc), "deletable": can_delete_document(user, doc)} for doc in docs]
    ctx = nav_context(request)
    ctx.update({"doc_cards": doc_cards})
    return templates.TemplateResponse(request, "document_trash.html", ctx)


@app.post("/documents/{document_id}/restore")
def restore_document(request: Request, document_id: int):
    user = require_user(request)
    require_write_access(request)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        item = get_document_for_user(conn, document_id, user)
        if not can_delete_document(user, item):
            raise HTTPException(status_code=403, detail="只有文档上传者本人或管理员可以恢复该文档")
        conn.execute("UPDATE document_files SET deleted_at='', deleted_by=NULL, updated_at=? WHERE id=?", (now, document_id))
        if item["task_id"]:
            conn.execute("UPDATE tasks SET content=?, updated_at=? WHERE id=?", (item["extracted_text"] or "", now, item["task_id"]))
            update_task_status(conn, item["task_id"])
    return RedirectResponse("/documents/trash", status_code=303)


@app.post("/documents/{document_id}/purge")
def purge_document(request: Request, document_id: int):
    user = require_user(request)
    require_write_access(request)
    with db_conn() as conn:
        item = get_document_for_user(conn, document_id, user)
        if not can_delete_document(user, item):
            raise HTTPException(status_code=403, detail="只有文档上传者本人或管理员可以永久删除该文档")
        delete_document_record(conn, document_id)
    return RedirectResponse("/documents/trash", status_code=303)


@app.post("/documents/{document_id}/ai/summarize")
def summarize_document(request: Request, document_id: int):
    user = require_user(request)
    require_write_access(request)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        item = get_document_for_user(conn, document_id, user)
        summary, provider = summarize_with_ai_api(item["extracted_text"] or "")
        conn.execute(
            """
            UPDATE document_files SET ai_summary=?, ai_summary_updated_at=?, ai_summary_provider=?
            WHERE id=?
            """,
            (summary, now, provider, document_id),
        )
    return RedirectResponse(f"/documents/{document_id}/edit#ai-summary", status_code=303)


@app.post("/tasks/{task_id}/images/upload")
def upload_task_image(request: Request, task_id: int, caption: str = Form(""), file: UploadFile = File(...)):
    user = require_user(request)
    require_write_access(request)
    suffix = Path(file.filename or "").suffix.lower()
    allowed = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
    if suffix not in allowed:
        raise HTTPException(status_code=400, detail="当前插图仅支持 png、jpg、jpeg、webp、gif、bmp 格式")
    with db_conn() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        if not task_accessible_for_images(user, task):
            raise HTTPException(status_code=403, detail="你没有权限给该任务上传插图")
        stored_name = safe_stored_name(suffix)
        dest = TASK_IMAGE_DIR / stored_name
        with dest.open("wb") as f:
            shutil.copyfileobj(file.file, f)
        mime_type = file.content_type or mimetypes.guess_type(file.filename or "")[0] or "application/octet-stream"
        conn.execute(
            """
            INSERT INTO task_images(task_id, original_name, stored_name, mime_type, caption, uploaded_by, uploaded_at)
            VALUES(?,?,?,?,?,?,?)
            """,
            (task_id, file.filename, stored_name, mime_type, caption.strip(), user["id"], datetime.now().isoformat(timespec="seconds")),
        )
    return RedirectResponse(f"/tasks/{task_id}#task-images", status_code=303)


@app.get("/tasks/{task_id}/images/{image_id}/view")
def view_task_image(request: Request, task_id: int, image_id: int):
    user = require_user(request)
    with db_conn() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        if user["role"] == "student" and task["assigned_to"] != user["id"]:
            raise HTTPException(status_code=403, detail="你只能查看自己任务中的插图")
        item = conn.execute("SELECT * FROM task_images WHERE id=? AND task_id=?", (image_id, task_id)).fetchone()
    if not item:
        raise HTTPException(status_code=404, detail="插图不存在")
    path = TASK_IMAGE_DIR / item["stored_name"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="插图文件丢失")
    return FileResponse(path, media_type=item["mime_type"] or "application/octet-stream", filename=item["original_name"])


@app.post("/tasks/{task_id}/images/{image_id}/delete")
def delete_task_image(request: Request, task_id: int, image_id: int):
    user = require_user(request)
    require_write_access(request)
    with db_conn() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        item = conn.execute("SELECT * FROM task_images WHERE id=? AND task_id=?", (image_id, task_id)).fetchone()
        if not task or not item:
            raise HTTPException(status_code=404, detail="插图不存在")
        if user["role"] not in {"teacher", "admin"} and item["uploaded_by"] != user["id"]:
            raise HTTPException(status_code=403, detail="只有上传者本人、教师或管理员可以删除该插图")
        cleanup_file(TASK_IMAGE_DIR / item["stored_name"])
        conn.execute("DELETE FROM task_images WHERE id=? AND task_id=?", (image_id, task_id))
    return RedirectResponse(f"/tasks/{task_id}#task-images", status_code=303)


@app.get("/templates", response_class=HTMLResponse)
def template_list(request: Request):
    user = require_user(request)
    with db_conn() as conn:
        files = conn.execute(
            """
            SELECT f.*, u.display_name AS uploader
            FROM template_files f JOIN users u ON f.uploaded_by=u.id
            ORDER BY f.uploaded_at DESC
            """
        ).fetchall()
    ctx = nav_context(request)
    ctx.update({"files": files})
    return templates.TemplateResponse(request, "templates.html", ctx)


@app.post("/templates/upload")
def upload_template(request: Request, description: str = Form(""), file: UploadFile = File(...)):
    user = require_user(request)
    require_write_access(request)
    require_teacher_or_admin(user)
    allowed = {".docx", ".pptx", ".xlsx", ".doc", ".ppt", ".pdf"}
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in allowed:
        raise HTTPException(status_code=400, detail="当前仅允许上传 docx、pptx、xlsx、doc、ppt、pdf 文件")
    stored_name = safe_stored_name(suffix)
    dest = TEMPLATE_FILE_DIR / stored_name
    with dest.open("wb") as f:
        shutil.copyfileobj(file.file, f)
    with db_conn() as conn:
        conn.execute(
            """
            INSERT INTO template_files(original_name, stored_name, uploaded_by, uploaded_at, description)
            VALUES(?,?,?,?,?)
            """,
            (file.filename, stored_name, user["id"], datetime.now().isoformat(timespec="seconds"), description.strip()),
        )
    return RedirectResponse("/templates", status_code=303)


@app.get("/templates/{file_id}/download")
def download_template(request: Request, file_id: int):
    require_user(request)
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM template_files WHERE id=?", (file_id,)).fetchone()
    if not item:
        raise HTTPException(status_code=404, detail="文件不存在")
    path = TEMPLATE_FILE_DIR / item["stored_name"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="文件丢失")
    return FileResponse(path, filename=item["original_name"], media_type="application/octet-stream")


@app.get("/admin/users", response_class=HTMLResponse)
def users_page(request: Request, error: str | None = None, message: str | None = None):
    user = require_user(request)
    require_admin(user)
    with db_conn() as conn:
        users = conn.execute(
            """
            SELECT u.*, COUNT(s.id) AS signin_days
            FROM users u LEFT JOIN signins s ON u.id=s.user_id AND s.signin_date LIKE ?
            GROUP BY u.id
            ORDER BY CASE u.role WHEN 'admin' THEN 1 WHEN 'teacher' THEN 2 ELSE 3 END, u.id
            """,
            (date.today().strftime("%Y-%m") + "%",),
        ).fetchall()
        request_rows = conn.execute(
            """
            SELECT r.*, applicant.display_name AS applicant_name, applicant.username AS applicant_username,
                   created.username AS created_username
            FROM admin_creation_requests r
            JOIN users applicant ON applicant.id=r.requested_by
            LEFT JOIN users created ON created.id=r.created_user_id
            ORDER BY CASE r.status WHEN '待审批' THEN 1 ELSE 2 END, r.id DESC
            """
        ).fetchall()
        admin_requests = []
        for row in request_rows:
            approvals = conn.execute(
                """
                SELECT a.admin_id, u.display_name, u.username
                FROM admin_creation_approvals a
                JOIN users u ON u.id=a.admin_id
                WHERE a.request_id=?
                ORDER BY a.id
                """,
                (row["id"],),
            ).fetchall()
            item = dict(row)
            item["approval_count"] = len(approvals)
            item["approved_by"] = "、".join(a["display_name"] for a in approvals)
            item["current_admin_approved"] = any(a["admin_id"] == user["id"] for a in approvals)
            item["can_approve"] = (
                row["status"] == "待审批"
                and row["requested_by"] != user["id"]
                and not item["current_admin_approved"]
            )
            item["can_withdraw"] = row["status"] == "待审批" and row["requested_by"] == user["id"]
            admin_requests.append(item)
    ctx = nav_context(request)
    ctx.update({"users": users, "admin_requests": admin_requests, "error": error, "message": message})
    return templates.TemplateResponse(request, "users.html", ctx)


@app.post("/admin/requests")
def create_admin_request(
    request: Request,
    display_name: str = Form(...),
    password: str = Form(...),
    password_confirm: str = Form(...),
    phone: str = Form(""),
    email: str = Form(""),
):
    user = require_user(request)
    require_write_access(request)
    require_admin(user)
    display_name = display_name.strip()
    if not display_name:
        return RedirectResponse(f"/admin/users?error={quote('姓名不能为空')}", status_code=303)
    if not password:
        return RedirectResponse(f"/admin/users?error={quote('密码不能为空')}", status_code=303)
    if password != password_confirm:
        return RedirectResponse(f"/admin/users?error={quote('两次输入的密码不一致')}", status_code=303)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        other_admins = conn.execute(
            "SELECT COUNT(*) FROM users WHERE role='admin' AND id<>?",
            (user["id"],),
        ).fetchone()[0]
        if other_admins < 2:
            return RedirectResponse(f"/admin/users?error={quote('当前可审批的其他管理员不足两名')}", status_code=303)
        username = next_admin_username(conn)
        conn.execute(
            """
            INSERT INTO admin_creation_requests(
                requested_by, username, display_name, password_hash, phone, email, status, required_approvals, created_at
            ) VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (user["id"], username, display_name, password_hash(password), phone.strip(), email.strip(), "待审批", 2, now),
        )
    return RedirectResponse(
        f"/admin/users?message={quote('管理员创建申请已提交，拟用登录账号为：' + username)}",
        status_code=303,
    )


@app.post("/admin/requests/{request_id}/approve")
def approve_admin_request(request: Request, request_id: int):
    admin_user = require_user(request)
    require_write_access(request)
    require_admin(admin_user)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM admin_creation_requests WHERE id=?", (request_id,)).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="申请不存在")
        if item["status"] != "待审批":
            return RedirectResponse(f"/admin/users?error={quote('该申请已结束')}", status_code=303)
        if item["requested_by"] == admin_user["id"]:
            return RedirectResponse(f"/admin/users?error={quote('申请人不能审批自己的申请')}", status_code=303)
        try:
            conn.execute(
                "INSERT INTO admin_creation_approvals(request_id, admin_id, created_at) VALUES(?,?,?)",
                (request_id, admin_user["id"], now),
            )
        except sqlite3.IntegrityError:
            return RedirectResponse(f"/admin/users?error={quote('你已经同意过该申请')}", status_code=303)

        approval_count = conn.execute(
            """
            SELECT COUNT(*)
            FROM admin_creation_approvals a
            JOIN users u ON u.id=a.admin_id
            WHERE a.request_id=? AND u.role='admin' AND a.admin_id<>?
            """,
            (request_id, item["requested_by"]),
        ).fetchone()[0]
        if approval_count >= item["required_approvals"]:
            if conn.execute("SELECT 1 FROM users WHERE username=?", (item["username"],)).fetchone():
                return RedirectResponse(f"/admin/users?error={quote('拟用登录账号已被占用')}", status_code=303)
            cursor = conn.execute(
                """
                INSERT INTO users(username, display_name, role, password_hash, student_id, phone, email, created_at, updated_at)
                VALUES(?,?,?,?,?,?,?,?,?)
                """,
                (item["username"], item["display_name"], "admin", item["password_hash"], None, item["phone"], item["email"], now, now),
            )
            conn.execute(
                """
                UPDATE admin_creation_requests
                SET status='已通过', resolved_at=?, created_user_id=?
                WHERE id=?
                """,
                (now, cursor.lastrowid, request_id),
            )
            return RedirectResponse(
                f"/admin/users?message={quote('两名管理员已同意，新管理员账号已创建：' + item['username'])}",
                status_code=303,
            )
    return RedirectResponse(f"/admin/users?message={quote('已同意该管理员创建申请')}", status_code=303)


@app.post("/admin/requests/{request_id}/withdraw")
def withdraw_admin_request(request: Request, request_id: int):
    admin_user = require_user(request)
    require_write_access(request)
    require_admin(admin_user)
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM admin_creation_requests WHERE id=?", (request_id,)).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="申请不存在")
        if item["requested_by"] != admin_user["id"]:
            return RedirectResponse(f"/admin/users?error={quote('只能撤回自己提交的申请')}", status_code=303)
        if item["status"] != "待审批":
            return RedirectResponse(f"/admin/users?error={quote('该申请已结束')}", status_code=303)
        conn.execute(
            "UPDATE admin_creation_requests SET status='已撤回', resolved_at=? WHERE id=?",
            (datetime.now().isoformat(timespec="seconds"), request_id),
        )
    return RedirectResponse(f"/admin/users?message={quote('管理员创建申请已撤回')}", status_code=303)


@app.post("/admin/teachers")
def create_teacher_user(
    request: Request,
    display_name: str = Form(...),
    password: str = Form(...),
    password_confirm: str = Form(...),
    phone: str = Form(""),
    email: str = Form(""),
):
    user = require_user(request)
    require_write_access(request)
    require_admin(user)
    display_name = display_name.strip()
    if not display_name:
        return RedirectResponse(f"/admin/users?error={quote('姓名不能为空')}", status_code=303)
    if not password:
        return RedirectResponse(f"/admin/users?error={quote('密码不能为空')}", status_code=303)
    if password != password_confirm:
        return RedirectResponse(f"/admin/users?error={quote('两次输入的密码不一致')}", status_code=303)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        username = next_teacher_username(conn)
        conn.execute(
            """
            INSERT INTO users(username, display_name, role, password_hash, student_id, phone, email, created_at, updated_at)
            VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (username, display_name, "teacher", password_hash(password), None, phone.strip(), email.strip(), now, now),
        )
    return RedirectResponse(f"/admin/users?message={quote('教师账号已创建，登录账号为：' + username)}", status_code=303)


@app.get("/admin/users/{target_user_id}/edit", response_class=HTMLResponse)
def edit_user_page(request: Request, target_user_id: int):
    user = require_user(request)
    require_admin(user)
    with db_conn() as conn:
        target = conn.execute("SELECT * FROM users WHERE id=?", (target_user_id,)).fetchone()
    if not target:
        raise HTTPException(status_code=404, detail="用户不存在")
    ctx = nav_context(request)
    ctx.update({"target": target, "error": None})
    return templates.TemplateResponse(request, "user_edit.html", ctx)


@app.post("/admin/users/{target_user_id}/edit")
def update_user(
    request: Request,
    target_user_id: int,
    display_name: str = Form(...),
    student_id: str = Form(""),
    phone: str = Form(""),
    email: str = Form(""),
    password: str = Form(""),
    password_confirm: str = Form(""),
):
    admin_user = require_user(request)
    require_write_access(request)
    require_admin(admin_user)
    display_name = display_name.strip()
    student_id = student_id.strip()
    phone = phone.strip()
    email = email.strip()
    with db_conn() as conn:
        target = conn.execute("SELECT * FROM users WHERE id=?", (target_user_id,)).fetchone()
        if not target:
            raise HTTPException(status_code=404, detail="用户不存在")
        values = {"display_name": display_name, "student_id": student_id, "phone": phone, "email": email}
        if not display_name:
            ctx = nav_context(request)
            ctx.update({"target": target, "error": "姓名不能为空", "values": values})
            return templates.TemplateResponse(request, "user_edit.html", ctx, status_code=400)
        new_username = target["username"]
        normalized_student_id = target["student_id"]
        if target["role"] == "student":
            if not student_id or not valid_student_id(student_id):
                ctx = nav_context(request)
                ctx.update({"target": target, "error": "学号格式不正确", "values": values})
                return templates.TemplateResponse(request, "user_edit.html", ctx, status_code=400)
            if count_users_by_student_id(conn, student_id, target_user_id) > 0 or conn.execute("SELECT 1 FROM users WHERE username=? AND id<>?", (student_id, target_user_id)).fetchone():
                ctx = nav_context(request)
                ctx.update({"target": target, "error": "该学号已被使用，请使用其他学号", "values": values})
                return templates.TemplateResponse(request, "user_edit.html", ctx, status_code=400)
            normalized_student_id = student_id
            new_username = student_id
        else:
            normalized_student_id = None
        if password or password_confirm:
            if password != password_confirm:
                ctx = nav_context(request)
                ctx.update({"target": target, "error": "两次输入的密码不一致", "values": values})
                return templates.TemplateResponse(request, "user_edit.html", ctx, status_code=400)
            conn.execute(
                """
                UPDATE users SET username=?, display_name=?, student_id=?, phone=?, email=?, password_hash=?, updated_at=?
                WHERE id=?
                """,
                (new_username, display_name, normalized_student_id, phone, email, password_hash(password), datetime.now().isoformat(timespec="seconds"), target_user_id),
            )
        else:
            conn.execute(
                """
                UPDATE users SET username=?, display_name=?, student_id=?, phone=?, email=?, updated_at=?
                WHERE id=?
                """,
                (new_username, display_name, normalized_student_id, phone, email, datetime.now().isoformat(timespec="seconds"), target_user_id),
            )
    return RedirectResponse("/admin/users", status_code=303)


@app.post("/admin/users/{target_user_id}/delete")
def delete_user(request: Request, target_user_id: int):
    admin_user = require_user(request)
    require_write_access(request)
    require_admin(admin_user)
    if target_user_id == admin_user["id"]:
        return RedirectResponse(f"/admin/users?error={quote('不能删除当前登录的管理员账号')}", status_code=303)
    with db_conn() as conn:
        target = conn.execute("SELECT * FROM users WHERE id=?", (target_user_id,)).fetchone()
        if not target:
            raise HTTPException(status_code=404, detail="用户不存在")
        if target["role"] == "admin":
            return RedirectResponse(f"/admin/users?error={quote('管理员账号不能通过页面删除')}", status_code=303)
        now = datetime.now().isoformat(timespec="seconds")
        conn.execute("DELETE FROM signins WHERE user_id=?", (target_user_id,))
        conn.execute("DELETE FROM personal_tasks WHERE owner_id=?", (target_user_id,))
        conn.execute("UPDATE tasks SET assigned_to=NULL, updated_at=? WHERE assigned_to=?", (now, target_user_id))
        conn.execute("UPDATE tasks SET created_by=? WHERE created_by=?", (admin_user["id"], target_user_id))
        conn.execute("UPDATE template_files SET uploaded_by=? WHERE uploaded_by=?", (admin_user["id"], target_user_id))
        conn.execute("UPDATE document_files SET created_by=? WHERE created_by=?", (admin_user["id"], target_user_id))
        conn.execute("UPDATE document_files SET updated_by=? WHERE updated_by=?", (admin_user["id"], target_user_id))
        conn.execute("UPDATE document_versions SET saved_by=? WHERE saved_by=?", (admin_user["id"], target_user_id))
        conn.execute("UPDATE task_images SET uploaded_by=? WHERE uploaded_by=?", (admin_user["id"], target_user_id))
        conn.execute("DELETE FROM users WHERE id=?", (target_user_id,))
    return RedirectResponse(f"/admin/users?message={quote('账号已删除，关联任务和文档记录已保留')}", status_code=303)


@app.get("/students/tasks", response_class=HTMLResponse)
def student_task_overview(request: Request):
    user = require_user(request)
    require_teacher_or_admin(user)
    with db_conn() as conn:
        students = conn.execute(
            """
            SELECT u.*, COUNT(s.id) AS signin_days
            FROM users u LEFT JOIN signins s ON u.id=s.user_id AND s.signin_date LIKE ?
            WHERE u.role='student'
            GROUP BY u.id
            ORDER BY u.display_name, u.id
            """,
            (date.today().strftime("%Y-%m") + "%",),
        ).fetchall()
        panels: list[dict[str, Any]] = []
        for student in students:
            generate_recurring_personal_tasks(student["id"], conn)
            assigned = conn.execute(
                """
                SELECT t.*, u.display_name AS assigned_name
                FROM tasks t LEFT JOIN users u ON t.assigned_to=u.id
                WHERE t.assigned_to=? AND t.status!='已完成'
                ORDER BY t.due_date IS NULL, t.due_date, t.id
                """,
                (student["id"],),
            ).fetchall()
            personal = conn.execute(
                """
                SELECT * FROM personal_tasks
                WHERE owner_id=? AND status!='已完成' AND COALESCE(is_template,0)=0
                ORDER BY due_date IS NULL, due_date, frequency, id
                """,
                (student["id"],),
            ).fetchall()
            panels.append(
                {
                    "student": student,
                    "assigned_cards": [{"task": t, "progress": task_progress_for_card(conn, t)} for t in assigned],
                    "personal_cards": [personal_task_view(t) for t in personal],
                    "unfinished_count": len(assigned) + len(personal),
                    "assigned_count": len(assigned),
                    "personal_count": len(personal),
                    "signin_calendar": month_signin_calendar(student["id"]),
                }
            )
    ctx = nav_context(request)
    ctx.update({"student_panels": panels})
    return templates.TemplateResponse(request, "student_tasks.html", ctx)


@app.get("/personal-tasks", response_class=HTMLResponse)
def personal_tasks_page(request: Request, error: str | None = None, message: str | None = None):
    user = require_user(request)
    require_student(user)
    with db_conn() as conn:
        generate_recurring_personal_tasks(user["id"], conn)
        template_rows = conn.execute(
            """
            SELECT * FROM personal_tasks
            WHERE owner_id=? AND COALESCE(is_template,0)=1
            ORDER BY auto_renew DESC, frequency, id DESC
            """,
            (user["id"],),
        ).fetchall()
        rows = conn.execute(
            """
            SELECT * FROM personal_tasks
            WHERE owner_id=? AND COALESCE(is_template,0)=0
            ORDER BY status='已完成', due_date IS NULL, due_date, frequency, id DESC
            """,
            (user["id"],),
        ).fetchall()
        personal_cards = []
        for row in rows:
            view = personal_task_view(row)
            view["block_progress"] = personal_once_block_progress(conn, row)
            personal_cards.append(view)
        template_cards = [personal_task_view(x) for x in template_rows]
    ctx = nav_context(request)
    ctx.update({"personal_cards": personal_cards, "template_cards": template_cards, "error": error, "message": message, "today": date.today().isoformat()})
    return templates.TemplateResponse(request, "personal_tasks.html", ctx)


@app.post("/personal-tasks/new")
def create_personal_task(
    request: Request,
    title: str = Form(...),
    description: str = Form(""),
    frequency: str = Form("daily"),
    due_date: str = Form(""),
):
    user = require_user(request)
    require_write_access(request)
    require_student(user)
    if frequency not in {"daily", "weekly", "once"}:
        raise HTTPException(status_code=400, detail="周期类型不正确")
    if not due_date:
        return RedirectResponse(f"/personal-tasks?error={quote('请为个人任务指定截止日期')}", status_code=303)
    start_date = date.today().isoformat()
    date_error = validate_due_date_not_before_start(due_date, start_date)
    if date_error:
        return RedirectResponse(f"/personal-tasks?error={quote(date_error)}", status_code=303)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        if frequency in {"daily", "weekly"}:
            conn.execute(
                """
                INSERT INTO personal_tasks(owner_id, title, description, frequency, start_date, status, due_date,
                                           created_at, updated_at, is_template, auto_renew)
                VALUES(?,?,?,?,?,?,?,?,?,?,?)
                """,
                (user["id"], title.strip(), description.strip(), frequency, start_date, "未完成", due_date, now, now, 1, 1),
            )
            generate_recurring_personal_tasks(user["id"], conn)
        else:
            conn.execute(
                """
                INSERT INTO personal_tasks(owner_id, title, description, frequency, start_date, status, due_date,
                                           created_at, updated_at, is_template, auto_renew)
                VALUES(?,?,?,?,?,?,?,?,?,?,?)
                """,
                (user["id"], title.strip(), description.strip(), frequency, start_date, "未完成", due_date, now, now, 0, 0),
            )
    return RedirectResponse("/personal-tasks", status_code=303)


@app.post("/personal-tasks/{task_id}/update")
def update_personal_task(
    request: Request,
    task_id: int,
    title: str = Form(...),
    description: str = Form(""),
    frequency: str = Form("daily"),
    due_date: str = Form(""),
):
    user = require_user(request)
    require_write_access(request)
    require_student(user)
    if frequency not in {"daily", "weekly", "once"}:
        raise HTTPException(status_code=400, detail="周期类型不正确")
    if not due_date:
        return RedirectResponse(f"/personal-tasks?error={quote('请为个人任务指定截止日期')}", status_code=303)
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM personal_tasks WHERE id=? AND owner_id=?", (task_id, user["id"])).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="个人任务不存在或无权修改")
        start_date = item["start_date"] or (item["created_at"] or date.today().isoformat())[:10]
        date_error = validate_due_date_not_before_start(due_date, start_date)
        if date_error:
            return RedirectResponse(f"/personal-tasks?error={quote(date_error)}", status_code=303)
        new_is_template = 1 if frequency in {"daily", "weekly"} and item["recurrence_parent_id"] is None else int(item["is_template"] or 0)
        new_auto = int(item["auto_renew"] or 0) if new_is_template else 0
        conn.execute(
            """
            UPDATE personal_tasks
            SET title=?, description=?, frequency=?, due_date=?, is_template=?, auto_renew=?, updated_at=?
            WHERE id=? AND owner_id=?
            """,
            (title.strip(), description.strip(), frequency, due_date, new_is_template, new_auto, datetime.now().isoformat(timespec="seconds"), task_id, user["id"]),
        )
        generate_recurring_personal_tasks(user["id"], conn)
    return RedirectResponse("/personal-tasks", status_code=303)


@app.post("/personal-tasks/{task_id}/cancel-renew")
def cancel_personal_task_renew(request: Request, task_id: int):
    user = require_user(request)
    require_write_access(request)
    require_student(user)
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM personal_tasks WHERE id=? AND owner_id=? AND COALESCE(is_template,0)=1", (task_id, user["id"])).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="自动续订规则不存在或无权修改")
        conn.execute("UPDATE personal_tasks SET auto_renew=0, updated_at=? WHERE id=? AND owner_id=?", (datetime.now().isoformat(timespec="seconds"), task_id, user["id"]))
    return RedirectResponse("/personal-tasks?message=" + quote("已取消自动续订"), status_code=303)


@app.post("/personal-tasks/{task_id}/resume-renew")
def resume_personal_task_renew(request: Request, task_id: int):
    user = require_user(request)
    require_write_access(request)
    require_student(user)
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM personal_tasks WHERE id=? AND owner_id=? AND COALESCE(is_template,0)=1", (task_id, user["id"])).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="自动续订规则不存在或无权修改")
        conn.execute("UPDATE personal_tasks SET auto_renew=1, updated_at=? WHERE id=? AND owner_id=?", (datetime.now().isoformat(timespec="seconds"), task_id, user["id"]))
        generate_recurring_personal_tasks(user["id"], conn)
    return RedirectResponse("/personal-tasks?message=" + quote("已恢复自动续订"), status_code=303)


@app.post("/personal-tasks/{task_id}/toggle")
def toggle_personal_task(request: Request, task_id: int):
    user = require_user(request)
    require_write_access(request)
    require_student(user)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM personal_tasks WHERE id=? AND owner_id=? AND COALESCE(is_template,0)=0", (task_id, user["id"])).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="个人任务不存在或无权修改")
        new_status = "未完成" if item["status"] == "已完成" else "已完成"
        completed_at = now if new_status == "已完成" else None
        conn.execute(
            "UPDATE personal_tasks SET status=?, completed_at=?, updated_at=? WHERE id=? AND owner_id=?",
            (new_status, completed_at, now, task_id, user["id"]),
        )
    return RedirectResponse("/personal-tasks", status_code=303)


@app.post("/personal-tasks/{task_id}/delete")
def delete_personal_task(request: Request, task_id: int):
    user = require_user(request)
    require_write_access(request)
    require_student(user)
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM personal_tasks WHERE id=? AND owner_id=?", (task_id, user["id"])).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="个人任务不存在或无权修改")
        conn.execute("DELETE FROM task_progress_records WHERE record_type='personal' AND task_id=?", (task_id,))
        if item["is_template"]:
            conn.execute("UPDATE personal_tasks SET recurrence_parent_id=NULL WHERE recurrence_parent_id=?", (task_id,))
        conn.execute("DELETE FROM personal_tasks WHERE id=? AND owner_id=?", (task_id, user["id"]))
    return RedirectResponse("/personal-tasks", status_code=303)


@app.post("/personal-tasks/{task_id}/complete-from-popup")
def complete_personal_task_from_popup(request: Request, task_id: int):
    user = require_user(request)
    require_write_access(request)
    require_student(user)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM personal_tasks WHERE id=? AND owner_id=?", (task_id, user["id"])).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="个人任务不存在或无权修改")
        conn.execute(
            "UPDATE personal_tasks SET status='已完成', completed_at=?, updated_at=? WHERE id=? AND owner_id=?",
            (now, now, task_id, user["id"]),
        )
    return RedirectResponse("/", status_code=303)



@app.post("/tasks/{task_id}/progress-records/new")
def create_assigned_progress_record(
    request: Request,
    task_id: int,
    block_index: int = Form(...),
    note: str = Form(""),
    file: UploadFile | None = File(None),
):
    user = require_user(request)
    require_write_access(request)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        if user["role"] != "student" or task["assigned_to"] != user["id"]:
            raise HTTPException(status_code=403, detail="只有任务负责人学生本人可以记录该任务进度")
        progress_info = assigned_task_block_progress(conn, task)
        blocks = progress_info["blocks"]
        if not progress_info["enabled"] or block_index < 0 or block_index >= len(blocks):
            raise HTTPException(status_code=400, detail="该任务当前不适用周区块进度记录")
        block = blocks[block_index]
        original_name, stored_name, mime_type = save_progress_upload(file)
        conn.execute(
            """
            INSERT INTO task_progress_records(record_type, task_id, owner_id, block_index, block_start, block_end,
                                              note, original_name, stored_name, mime_type, created_at)
            VALUES(?,?,?,?,?,?,?,?,?,?,?)
            """,
            ("assigned", task_id, user["id"], block_index, block["start"], block["end"], note.strip(), original_name, stored_name, mime_type, now),
        )
        update_task_status(conn, task_id)
    return RedirectResponse(f"/tasks/{task_id}#block-progress", status_code=303)


@app.post("/personal-tasks/{task_id}/progress-records/new")
def create_personal_progress_record(
    request: Request,
    task_id: int,
    block_index: int = Form(...),
    note: str = Form(""),
    file: UploadFile | None = File(None),
):
    user = require_user(request)
    require_write_access(request)
    require_student(user)
    now = datetime.now().isoformat(timespec="seconds")
    with db_conn() as conn:
        item = conn.execute("SELECT * FROM personal_tasks WHERE id=? AND owner_id=? AND frequency='once' AND COALESCE(is_template,0)=0", (task_id, user["id"])).fetchone()
        if not item:
            raise HTTPException(status_code=404, detail="个人一次性任务不存在或无权记录")
        progress_info = personal_once_block_progress(conn, item)
        blocks = progress_info["blocks"]
        if not progress_info["enabled"] or block_index < 0 or block_index >= len(blocks):
            raise HTTPException(status_code=400, detail="该个人任务当前不适用周区块进度记录")
        block = blocks[block_index]
        if not block.get("is_open"):
            return RedirectResponse(f"/personal-tasks?error={quote('该进展节点尚未开放或已超过可填写日期')}#personal-progress", status_code=303)
        if block.get("recorded"):
            return RedirectResponse(f"/personal-tasks?error={quote('该节点已经保存过本周进展记录')}#personal-progress", status_code=303)
        if not note.strip():
            return RedirectResponse(f"/personal-tasks?error={quote('本周进展文字记录不能为空')}#personal-progress", status_code=303)
        original_name, stored_name, mime_type = save_progress_upload(file)
        conn.execute(
            """
            INSERT INTO task_progress_records(record_type, task_id, owner_id, block_index, block_start, block_end,
                                              note, original_name, stored_name, mime_type, created_at)
            VALUES(?,?,?,?,?,?,?,?,?,?,?)
            """,
            ("personal", task_id, user["id"], block_index, block["start"], block["end"], note.strip(), original_name, stored_name, mime_type, now),
        )
        updated_progress = personal_once_block_progress(conn, item)
        if updated_progress["enabled"] and updated_progress["percent"] >= 100:
            conn.execute("UPDATE personal_tasks SET status='已完成', completed_at=?, updated_at=? WHERE id=? AND owner_id=?", (now, now, task_id, user["id"]))
    return RedirectResponse("/personal-tasks#personal-progress", status_code=303)


@app.get("/progress-records/{record_id}/download")
def download_progress_record(request: Request, record_id: int):
    user = require_user(request)
    with db_conn() as conn:
        record = conn.execute("SELECT * FROM task_progress_records WHERE id=?", (record_id,)).fetchone()
        if not record:
            raise HTTPException(status_code=404, detail="进度记录不存在")
        if not progress_record_access(conn, record, user):
            raise HTTPException(status_code=403, detail="你没有权限下载该进度记录附件")
    if not record["stored_name"]:
        raise HTTPException(status_code=404, detail="该进度记录没有附件")
    path = PROJECT_FILE_DIR / record["stored_name"]
    if not path.exists():
        raise HTTPException(status_code=404, detail="附件文件丢失")
    return FileResponse(path, media_type=record["mime_type"] or "application/octet-stream", filename=record["original_name"] or "progress_attachment")


@app.get("/ai", response_class=HTMLResponse)
def ai_placeholder(request: Request):
    user = require_user(request)
    with db_conn() as conn:
        docs = conn.execute(
            """
            SELECT d.*, t.title AS task_title, t.parent_id, parent.title AS parent_title, u.display_name AS creator_name
            FROM document_files d
            LEFT JOIN tasks t ON d.task_id=t.id
            LEFT JOIN tasks parent ON t.parent_id=parent.id
            LEFT JOIN users u ON d.created_by=u.id
            WHERE d.deleted_at IS NULL OR d.deleted_at=''
            ORDER BY d.updated_at DESC, d.id DESC
            """
        ).fetchall()
        doc_cards = []
        for doc in docs:
            doc_cards.append({
                "doc": doc,
                "word_count": document_word_count(doc),
                "warnings": document_ai_warnings(conn, doc["id"]),
            })
    ctx = nav_context(request)
    ctx.update({
        "doc_cards": doc_cards,
        "ai_configured": bool(os.getenv("AI_API_URL", "").strip() and os.getenv("AI_API_KEY", "").strip()),
        "ai_model": os.getenv("AI_MODEL", ""),
    })
    return templates.TemplateResponse(request, "ai_placeholder.html", ctx)


@app.get("/demo", response_class=HTMLResponse)
def demo_preview(request: Request):
    """无需登录的界面演示页，便于在服务器部署前预览视觉效果。"""
    return templates.TemplateResponse(request, "demo_preview.html", {"request": request, "user": None, "title": "界面 Demo 预览"})
